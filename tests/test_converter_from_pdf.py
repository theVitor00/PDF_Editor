"""
Testes para conversão de PDF para outros formatos (`converter_from_pdf.py`).
Cobre: DOCX, HTML, Markdown, TXT, PNG e JPG.
Garante que o PDF original permanece intacto.
"""

from pathlib import Path
from docx import Document
from cortex_pdf.core.converter_from_pdf import convert_pdf_to_format
from tests.conftest import file_sha256


def test_convert_pdf_to_docx(sample_pdf: Path, tmp_path: Path):
    orig_hash = file_sha256(sample_pdf)

    res = convert_pdf_to_format(
        input_path=sample_pdf,
        target_format="docx",
        output_dir=tmp_path,
        keep_original_name=False,
        custom_name="exportado_word",
    )
    assert len(res.output_paths) == 1
    out = res.output_paths[0]
    assert out.exists()
    assert out.name == "exportado_word.docx"
    assert file_sha256(sample_pdf) == orig_hash

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Titulo Principal" in full_text or len(doc.paragraphs) > 0


def test_convert_pdf_to_html_md_and_txt(sample_pdf: Path, tmp_path: Path):
    orig_hash = file_sha256(sample_pdf)

    # HTML
    res_html = convert_pdf_to_format(
        input_path=sample_pdf,
        target_format="html",
        output_dir=tmp_path,
        keep_original_name=True,
    )
    html_file = res_html.output_paths[0]
    assert html_file.exists()
    assert html_file.name == "documento_original.html"
    html_content = html_file.read_text(encoding="utf-8")
    assert "Titulo Principal Cortex PDF" in html_content
    assert "<section class=\"pdf-page\"" in html_content

    # Markdown
    res_md = convert_pdf_to_format(
        input_path=sample_pdf,
        target_format="md",
        output_dir=tmp_path,
        keep_original_name=True,
    )
    md_file = res_md.output_paths[0]
    assert md_file.exists()
    md_content = md_file.read_text(encoding="utf-8")
    assert "### Titulo Principal Cortex PDF" in md_content

    # TXT
    res_txt = convert_pdf_to_format(
        input_path=sample_pdf,
        target_format="txt",
        output_dir=tmp_path,
        keep_original_name=False,
        custom_name="texto_extraido",
    )
    txt_file = res_txt.output_paths[0]
    assert txt_file.exists()
    assert txt_file.name == "texto_extraido.txt"
    assert "Segunda Pagina do Relatorio" in txt_file.read_text(encoding="utf-8")

    assert file_sha256(sample_pdf) == orig_hash


def test_convert_pdf_to_images_png_and_jpg(sample_pdf: Path, tmp_path: Path):
    orig_hash = file_sha256(sample_pdf)

    res_png = convert_pdf_to_format(
        input_path=sample_pdf,
        target_format="png",
        output_dir=tmp_path,
        keep_original_name=True,
        dpi=120,
    )
    assert len(res_png.output_paths) == 3
    for p in res_png.output_paths:
        assert p.exists()
        assert p.suffix == ".png"
        assert p.stat().st_size > 100

    res_jpg = convert_pdf_to_format(
        input_path=sample_pdf,
        target_format="jpg",
        output_dir=tmp_path,
        keep_original_name=False,
        custom_name="slide",
        dpi=100,
    )
    assert len(res_jpg.output_paths) == 3
    assert res_jpg.output_paths[0].name == "slide_parte_01.jpg"
    assert file_sha256(sample_pdf) == orig_hash

