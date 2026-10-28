"""
Fixtures compartilhadas para os testes de todas as funcionalidades do Cortex PDF Editor.
Inclui verificação de hash SHA-256 para comprovar que o arquivo original nunca é alterado.
"""

import hashlib
import os
from pathlib import Path
import openpyxl
from PIL import Image
import pymupdf
import pytest
from docx import Document

# Garante execução headless (offscreen) para testes PyQt6
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def file_sha256(path: Path) -> str:
    """Calcula o hash SHA-256 de um arquivo para garantir que não sofreu mutação."""
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """Cria um PDF de 3 páginas com textos variados para testes."""
    pdf_path = tmp_path / "documento_original.pdf"
    with pymupdf.open() as doc:
        p1 = doc.new_page(width=595, height=842)
        p1.insert_text(pymupdf.Point(72, 100), "Titulo Principal Cortex PDF", fontsize=20)
        p1.insert_text(
            pymupdf.Point(72, 150),
            "Este eh o texto da primeira pagina para selecao e edicao.",
            fontsize=12,
        )
        p1.insert_text(pymupdf.Point(72, 200), "Valor Antigo: 1000", fontsize=12)

        p2 = doc.new_page(width=595, height=842)
        p2.insert_text(pymupdf.Point(72, 100), "Segunda Pagina do Relatorio", fontsize=16)
        p2.insert_text(pymupdf.Point(72, 150), "Conteudo exclusivo da pagina dois.", fontsize=12)

        p3 = doc.new_page(width=595, height=842)
        p3.insert_text(pymupdf.Point(72, 100), "Terceira Pagina Final", fontsize=16)
        p3.insert_text(pymupdf.Point(72, 150), "Conclusao do documento de teste.", fontsize=12)

        doc.save(str(pdf_path))
    return pdf_path


@pytest.fixture
def second_pdf(tmp_path: Path) -> Path:
    """Cria um segundo PDF de 2 páginas para testes de junção (merge)."""
    pdf_path = tmp_path / "anexo_extra.pdf"
    with pymupdf.open() as doc:
        p1 = doc.new_page(width=595, height=842)
        p1.insert_text(pymupdf.Point(72, 100), "Anexo Parte A", fontsize=16)
        p2 = doc.new_page(width=595, height=842)
        p2.insert_text(pymupdf.Point(72, 100), "Anexo Parte B", fontsize=16)
        doc.save(str(pdf_path))
    return pdf_path


@pytest.fixture
def bloated_pdf(tmp_path: Path) -> Path:
    """Cria um PDF não-comprimido contendo imagem grande e fluxos sem deflate."""
    img_path = tmp_path / "raw_banner.png"
    img = Image.new("RGB", (1200, 900), color=(40, 120, 200))
    # Desenha gradiente para ter entropia realista
    pixels = img.load()
    for x in range(0, 1200, 4):
        for y in range(0, 900, 4):
            pixels[x, y] = ((x * 3) % 255, (y * 5) % 255, 180)
    img.save(str(img_path), format="PNG", compress_level=0)

    pdf_path = tmp_path / "pesado_original.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page(width=595, height=842)
        page.insert_text(pymupdf.Point(50, 50), "PDF com imagem pesada para compressao", fontsize=14)
        page.insert_image(pymupdf.Rect(50, 80, 500, 450), filename=str(img_path))
        # Salva sem deflate e sem garbage collection para permitir compressão mensurável
        doc.save(str(pdf_path), garbage=0, deflate=False)
    return pdf_path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Cria um arquivo Word (.docx) de exemplo."""
    path = tmp_path / "contrato_exemplo.docx"
    doc = Document()
    doc.add_heading("Contrato de Prestacao de Servicos", level=1)
    p = doc.add_paragraph("Este documento foi gerado em formato ")
    run = p.add_run("DOCX com negrito")
    run.bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Descricao"
    table.cell(1, 0).text = "01"
    table.cell(1, 1).text = "Consultoria Tecnica"
    doc.save(str(path))
    return path


@pytest.fixture
def sample_image(tmp_path: Path) -> Path:
    """Cria uma imagem PNG de exemplo."""
    path = tmp_path / "grafico_vendas.png"
    img = Image.new("RGB", (400, 300), color=(25, 135, 84))
    img.save(str(path))
    return path


@pytest.fixture
def sample_txt(tmp_path: Path) -> Path:
    """Cria um arquivo TXT de exemplo."""
    path = tmp_path / "notas_reuniao.txt"
    path.write_text("Ata de Reuniao\nLinha 1: Aprovacao do projeto.\nLinha 2: Prazo definido.", encoding="utf-8")
    return path


@pytest.fixture
def sample_md(tmp_path: Path) -> Path:
    """Cria um arquivo Markdown (.md) de exemplo."""
    path = tmp_path / "documentacao.md"
    path.write_text("# Guia Rapido\n## Introducao\n- Item A\n- Item B\nTexto final.", encoding="utf-8")
    return path


@pytest.fixture
def sample_html(tmp_path: Path) -> Path:
    """Cria um arquivo HTML de exemplo."""
    path = tmp_path / "pagina_web.html"
    path.write_text(
        "<html><body><h1>Relatorio HTML</h1><p>Paragrafo formatado para PDF.</p></body></html>",
        encoding="utf-8",
    )
    return path


@pytest.fixture
def sample_csv(tmp_path: Path) -> Path:
    """Cria um arquivo CSV de exemplo."""
    path = tmp_path / "dados_financeiros.csv"
    path.write_text("Mes,Receita,Despesa\nJaneiro,5000,3200\nFevereiro,6100,2900\n", encoding="utf-8")
    return path


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    """Cria uma planilha Excel (.xlsx) de exemplo."""
    path = tmp_path / "planilha_controle.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Codigo", "Produto", "Quantidade"])
    ws.append(["A10", "Licenca Cortex", 15])
    wb.save(str(path))
    wb.close()
    return path

