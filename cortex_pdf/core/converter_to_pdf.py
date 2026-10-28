"""
Motor de Conversão de Outros Formatos para PDF.
Suporta:
- Imagens (.png, .jpg, .jpeg, .bmp, .webp, .tiff)
- Documentos Word (.docx)
- Páginas HTML (.html, .htm)
- Arquivos de Texto e Markdown (.txt, .md)
- Planilhas e Tabelas (.xlsx, .csv)

O arquivo original de entrada nunca é modificado; um novo PDF é gerado.
"""

import csv
import html
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Union
import pymupdf

from cortex_pdf.core.output_naming import resolve_output_path


SUPPORTED_TO_PDF_EXTENSIONS = (
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".webp",
    ".tiff",
    ".docx",
    ".html",
    ".htm",
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
)


@dataclass
class ConversionToPDFResult:
    """Resultado da conversão de um arquivo externo para PDF."""
    input_path: Path
    output_path: Path
    source_format: str
    message: str


def convert_to_pdf(
    input_path: Union[str, Path],
    output_dir: Optional[Union[str, Path]] = None,
    keep_original_name: bool = True,
    custom_name: Optional[str] = None,
) -> ConversionToPDFResult:
    """
    Converte um arquivo de diversos formatos suportados para um documento PDF.
    """
    src = Path(input_path).resolve()
    if not src.exists():
        raise FileNotFoundError(f"Arquivo de entrada não encontrado: {src}")

    ext = src.suffix.lower()
    if ext not in SUPPORTED_TO_PDF_EXTENSIONS:
        raise ValueError(
            f"Extensão '{ext}' não suportada para conversão em PDF. "
            f"Formatos aceitos: {', '.join(SUPPORTED_TO_PDF_EXTENSIONS)}"
        )

    out_path = resolve_output_path(
        input_path=src,
        output_dir=output_dir,
        keep_original_name=keep_original_name,
        custom_name=custom_name,
        target_extension=".pdf",
        default_suffix="_convertido",
    )

    if ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tiff"):
        _image_to_pdf(src, out_path)
    elif ext == ".docx":
        _docx_to_pdf(src, out_path)
    elif ext in (".html", ".htm"):
        _html_to_pdf(src, out_path)
    elif ext in (".txt", ".md"):
        _text_or_md_to_pdf(src, out_path, is_markdown=(ext == ".md"))
    elif ext == ".csv":
        _csv_to_pdf(src, out_path)
    elif ext == ".xlsx":
        _xlsx_to_pdf(src, out_path)

    return ConversionToPDFResult(
        input_path=src,
        output_path=out_path,
        source_format=ext.lstrip("."),
        message=f"Arquivo '{src.name}' convertido para PDF com sucesso ({out_path.name}).",
    )


def _image_to_pdf(src: Path, out_path: Path) -> None:
    """Converte uma imagem para página PDF mantendo a proporção e resolução."""
    from PIL import Image

    with Image.open(src) as img:
        rgb_img = img.convert("RGB")
        width, height = rgb_img.size

    with pymupdf.open() as pdf_doc:
        page = pdf_doc.new_page(width=float(width), height=float(height))
        page.insert_image(pymupdf.Rect(0, 0, float(width), float(height)), filename=str(src))
        pdf_doc.save(str(out_path), garbage=2, deflate=True)


def _render_html_string_to_pdf(html_content: str, out_path: Path) -> None:
    """Renderiza HTML/CSS estruturado em páginas PDF A4 utilizando PyMuPDF Story."""
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)  # Margens de 0.5 polegada (36pt)

    story = pymupdf.Story(html=html_content)
    writer = pymupdf.DocumentWriter(str(out_path))
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()


def _docx_to_pdf(src: Path, out_path: Path) -> None:
    """Converte documento Word (.docx) em PDF preservando estilos de títulos, parágrafos e tabelas."""
    from docx import Document

    doc = Document(str(src))
    body_parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name if para.style else "").lower()
        escaped_runs = []
        for run in para.runs:
            r_text = html.escape(run.text)
            if not r_text:
                continue
            if run.bold:
                r_text = f"<b>{r_text}</b>"
            if run.italic:
                r_text = f"<i>{r_text}</i>"
            if run.underline:
                r_text = f"<u>{r_text}</u>"
            escaped_runs.append(r_text)

        inner_html = "".join(escaped_runs) if escaped_runs else html.escape(text)
        if "heading 1" in style_name or "título 1" in style_name:
            body_parts.append(f"<h1>{inner_html}</h1>")
        elif "heading 2" in style_name or "título 2" in style_name:
            body_parts.append(f"<h2>{inner_html}</h2>")
        elif "heading 3" in style_name or "título 3" in style_name:
            body_parts.append(f"<h3>{inner_html}</h3>")
        else:
            body_parts.append(f"<p>{inner_html}</p>")

    for table in doc.tables:
        rows_html = []
        for r_idx, row in enumerate(table.rows):
            cells_html = []
            tag = "th" if r_idx == 0 else "td"
            for cell in row.cells:
                cells_html.append(f"<{tag}>{html.escape(cell.text.strip())}</{tag}>")
            rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
        body_parts.append(f'<table border="1" cellspacing="0" cellpadding="4">{"".join(rows_html)}</table>')

    if not body_parts:
        body_parts.append("<p> </p>")

    full_html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        "<style>"
        "body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; color: #111827; }"
        "h1 { font-size: 20pt; color: #0f172a; margin-bottom: 8pt; }"
        "h2 { font-size: 15pt; color: #1e293b; margin-bottom: 6pt; }"
        "h3 { font-size: 13pt; color: #334155; margin-bottom: 4pt; }"
        "p { margin-bottom: 8pt; }"
        "table { width: 100%; border-collapse: collapse; margin: 10pt 0; }"
        "th, td { border: 1px solid #94a3b8; padding: 5pt; font-size: 10pt; }"
        "</style></head><body>"
        + "\n".join(body_parts)
        + "</body></html>"
    )
    _render_html_string_to_pdf(full_html, out_path)


def _html_to_pdf(src: Path, out_path: Path) -> None:
    """Converte arquivo HTML para PDF."""
    raw_html = src.read_text(encoding="utf-8", errors="replace")
    _render_html_string_to_pdf(raw_html, out_path)


def _text_or_md_to_pdf(src: Path, out_path: Path, is_markdown: bool = False) -> None:
    """Converte arquivo TXT ou Markdown para PDF."""
    raw_text = src.read_text(encoding="utf-8", errors="replace")
    lines = raw_text.splitlines()
    html_lines: list[str] = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            html_lines.append("<br/>")
            continue
        if is_markdown:
            if stripped.startswith("### "):
                html_lines.append(f"<h3>{html.escape(stripped[4:])}</h3>")
                continue
            elif stripped.startswith("## "):
                html_lines.append(f"<h2>{html.escape(stripped[3:])}</h2>")
                continue
            elif stripped.startswith("# "):
                html_lines.append(f"<h1>{html.escape(stripped[2:])}</h1>")
                continue
            elif stripped.startswith("- ") or stripped.startswith("* "):
                html_lines.append(f"<p>• {html.escape(stripped[2:])}</p>")
                continue
        html_lines.append(f"<p>{html.escape(line)}</p>")

    if not html_lines:
        html_lines.append("<p> </p>")

    full_html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        "<style>"
        "body { font-family: sans-serif; font-size: 11pt; line-height: 1.4; color: #111827; }"
        "p { margin: 2pt 0; }"
        "</style></head><body>"
        + "\n".join(html_lines)
        + "</body></html>"
    )
    _render_html_string_to_pdf(full_html, out_path)


def _csv_to_pdf(src: Path, out_path: Path) -> None:
    """Converte planilha CSV em tabela PDF formatada."""
    with src.open("r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    rows_html = []
    for r_idx, row in enumerate(rows):
        tag = "th" if r_idx == 0 else "td"
        cells = "".join(f"<{tag}>{html.escape(str(c))}</{tag}>" for c in row)
        rows_html.append(f"<tr>{cells}</tr>")

    table_html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        "<style>"
        "body { font-family: sans-serif; font-size: 10pt; }"
        "table { width: 100%; border-collapse: collapse; }"
        "th, td { border: 1px solid #64748b; padding: 4pt; }"
        "</style></head><body>"
        f"<h2>{html.escape(src.stem)}</h2>"
        f"<table border='1' cellspacing='0' cellpadding='4'>{''.join(rows_html)}</table>"
        "</body></html>"
    )
    _render_html_string_to_pdf(table_html, out_path)


def _xlsx_to_pdf(src: Path, out_path: Path) -> None:
    """Converte planilha Excel (.xlsx) em documento PDF."""
    import openpyxl

    wb = openpyxl.load_workbook(str(src), data_only=True)
    sections_html = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_html = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if all(v is None for v in row):
                continue
            tag = "th" if r_idx == 0 else "td"
            cells = "".join(
                f"<{tag}>{html.escape('' if v is None else str(v))}</{tag}>" for v in row
            )
            rows_html.append(f"<tr>{cells}</tr>")
        if rows_html:
            sections_html.append(
                f"<h2>{html.escape(sheet_name)}</h2>"
                f"<table border='1' cellspacing='0' cellpadding='4'>{''.join(rows_html)}</table>"
            )
    wb.close()

    if not sections_html:
        sections_html.append("<p>Planilha vazia</p>")

    full_html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        "<style>"
        "body { font-family: sans-serif; font-size: 10pt; }"
        "table { width: 100%; border-collapse: collapse; margin-bottom: 14pt; }"
        "th, td { border: 1px solid #64748b; padding: 4pt; }"
        "</style></head><body>"
        + "\n".join(sections_html)
        + "</body></html>"
    )
    _render_html_string_to_pdf(full_html, out_path)
