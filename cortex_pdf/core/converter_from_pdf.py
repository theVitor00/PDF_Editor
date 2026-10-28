"""
Motor de Conversão de PDF para diversos formatos:
- DOCX (Word) mantendo layout, tabelas, fontes e imagens (via pdf2docx)
- HTML mantendo formatação visual, fontes e posicionamento (via PyMuPDF HTML)
- Markdown (.md) estruturado
- Texto (.txt) preservando ordem de leitura de blocos
- Imagens (.png, .jpg) em alta resolução

Em todos os casos, o arquivo PDF original nunca é modificado.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Union
import pymupdf

from cortex_pdf.core.output_naming import resolve_output_path, resolve_multi_output_paths


SUPPORTED_FROM_PDF_FORMATS = ("docx", "html", "md", "txt", "png", "jpg")


@dataclass
class ConversionFromPDFResult:
    """Resultado da conversão de um PDF para outro formato."""
    input_path: Path
    output_paths: list[Path]
    target_format: str
    message: str


def convert_pdf_to_format(
    input_path: Union[str, Path],
    target_format: str,
    output_dir: Optional[Union[str, Path]] = None,
    keep_original_name: bool = True,
    custom_name: Optional[str] = None,
    dpi: int = 200,
) -> ConversionFromPDFResult:
    """
    Converte um arquivo PDF para o formato especificado mantendo o máximo de fidelidade.
    """
    src = Path(input_path).resolve()
    if not src.exists():
        raise FileNotFoundError(f"Arquivo PDF não encontrado: {src}")

    fmt = target_format.lower().strip().lstrip(".")
    if fmt not in SUPPORTED_FROM_PDF_FORMATS:
        raise ValueError(
            f"Formato de destino não suportado: '{fmt}'. "
            f"Suportados: {', '.join(SUPPORTED_FROM_PDF_FORMATS)}"
        )

    if fmt == "docx":
        return _convert_pdf_to_docx(src, output_dir, keep_original_name, custom_name)
    elif fmt == "html":
        return _convert_pdf_to_html(src, output_dir, keep_original_name, custom_name)
    elif fmt == "md":
        return _convert_pdf_to_markdown(src, output_dir, keep_original_name, custom_name)
    elif fmt == "txt":
        return _convert_pdf_to_txt(src, output_dir, keep_original_name, custom_name)
    elif fmt in ("png", "jpg"):
        return _convert_pdf_to_images(src, fmt, output_dir, keep_original_name, custom_name, dpi=dpi)

    raise ValueError(f"Formato inválido: {fmt}")


def _convert_pdf_to_docx(
    src: Path,
    output_dir: Optional[Union[str, Path]],
    keep_original_name: bool,
    custom_name: Optional[str],
) -> ConversionFromPDFResult:
    out_path = resolve_output_path(
        input_path=src,
        output_dir=output_dir,
        keep_original_name=keep_original_name,
        custom_name=custom_name,
        target_extension=".docx",
        default_suffix="_convertido",
    )

    try:
        from pdf2docx import Converter

        cv = Converter(str(src))
        try:
            cv.convert(str(out_path))
        finally:
            cv.close()
    except Exception:
        # Fallback seguro com python-docx preservando parágrafos e tamanhos de fonte por página
        from docx import Document
        from docx.shared import Pt

        word_doc = Document()
        with pymupdf.open(str(src)) as pdf_doc:
            for p_idx, page in enumerate(pdf_doc):
                if p_idx > 0:
                    word_doc.add_page_break()
                text_dict = page.get_text("dict")
                for block in text_dict.get("blocks", []):
                    if block.get("type") != 0:
                        continue
                    para = word_doc.add_paragraph()
                    for line in block.get("lines", []):
                        for span in line.get("spans", []):
                            txt = span.get("text", "")
                            if txt:
                                run = para.add_run(txt + " ")
                                size = float(span.get("size", 11.0))
                                run.font.size = Pt(max(8.0, min(36.0, size)))
        word_doc.save(str(out_path))

    return ConversionFromPDFResult(
        input_path=src,
        output_paths=[out_path],
        target_format="docx",
        message=f"PDF convertido para DOCX com sucesso: {out_path.name}",
    )


def _convert_pdf_to_html(
    src: Path,
    output_dir: Optional[Union[str, Path]],
    keep_original_name: bool,
    custom_name: Optional[str],
) -> ConversionFromPDFResult:
    out_path = resolve_output_path(
        input_path=src,
        output_dir=output_dir,
        keep_original_name=keep_original_name,
        custom_name=custom_name,
        target_extension=".html",
        default_suffix="_convertido",
    )

    pages_html = []
    with pymupdf.open(str(src)) as doc:
        for idx, page in enumerate(doc):
            page_html = page.get_text("html")
            pages_html.append(
                f'<section class="pdf-page" data-page="{idx + 1}">\n{page_html}\n</section>'
            )

    full_html = (
        "<!DOCTYPE html>\n"
        '<html lang="pt-BR">\n'
        "<head>\n"
        '  <meta charset="utf-8">\n'
        f"  <title>{src.stem}</title>\n"
        "  <style>\n"
        "    body { background: #f1f5f9; margin: 24px; font-family: sans-serif; }\n"
        "    .pdf-page { background: #ffffff; margin: 0 auto 24px auto; "
        "box-shadow: 0 4px 14px rgba(0,0,0,0.12); padding: 12px; position: relative; overflow: hidden; }\n"
        "  </style>\n"
        "</head>\n"
        "<body>\n"
        + "\n".join(pages_html)
        + "\n</body>\n</html>\n"
    )
    out_path.write_text(full_html, encoding="utf-8")

    return ConversionFromPDFResult(
        input_path=src,
        output_paths=[out_path],
        target_format="html",
        message=f"PDF convertido para HTML mantendo layout: {out_path.name}",
    )


def _convert_pdf_to_markdown(
    src: Path,
    output_dir: Optional[Union[str, Path]],
    keep_original_name: bool,
    custom_name: Optional[str],
) -> ConversionFromPDFResult:
    out_path = resolve_output_path(
        input_path=src,
        output_dir=output_dir,
        keep_original_name=keep_original_name,
        custom_name=custom_name,
        target_extension=".md",
        default_suffix="_convertido",
    )

    md_sections: list[str] = []
    with pymupdf.open(str(src)) as doc:
        for idx, page in enumerate(doc):
            md_sections.append(f"## Página {idx + 1}\n")
            text_dict = page.get_text("dict")
            for block in text_dict.get("blocks", []):
                if block.get("type") != 0:
                    continue
                lines_str = []
                max_font = 0.0
                for line in block.get("lines", []):
                    span_texts = []
                    for span in line.get("spans", []):
                        t = span.get("text", "").strip()
                        if t:
                            span_texts.append(t)
                            max_font = max(max_font, float(span.get("size", 11.0)))
                    if span_texts:
                        lines_str.append(" ".join(span_texts))
                block_text = "\n".join(lines_str).strip()
                if not block_text:
                    continue
                if max_font >= 18.0:
                    md_sections.append(f"### {block_text}\n")
                else:
                    md_sections.append(f"{block_text}\n")

    out_path.write_text("\n".join(md_sections), encoding="utf-8")

    return ConversionFromPDFResult(
        input_path=src,
        output_paths=[out_path],
        target_format="md",
        message=f"PDF convertido para Markdown: {out_path.name}",
    )


def _convert_pdf_to_txt(
    src: Path,
    output_dir: Optional[Union[str, Path]],
    keep_original_name: bool,
    custom_name: Optional[str],
) -> ConversionFromPDFResult:
    out_path = resolve_output_path(
        input_path=src,
        output_dir=output_dir,
        keep_original_name=keep_original_name,
        custom_name=custom_name,
        target_extension=".txt",
        default_suffix="_convertido",
    )

    pages_txt: list[str] = []
    with pymupdf.open(str(src)) as doc:
        for idx, page in enumerate(doc):
            txt = page.get_text("text", sort=True)
            pages_txt.append(f"--- Página {idx + 1} ---\n{txt.strip()}\n")

    out_path.write_text("\n".join(pages_txt), encoding="utf-8")

    return ConversionFromPDFResult(
        input_path=src,
        output_paths=[out_path],
        target_format="txt",
        message=f"PDF convertido para TXT: {out_path.name}",
    )


def _convert_pdf_to_images(
    src: Path,
    fmt: str,
    output_dir: Optional[Union[str, Path]],
    keep_original_name: bool,
    custom_name: Optional[str],
    dpi: int = 200,
) -> ConversionFromPDFResult:
    ext = f".{fmt}"
    zoom = max(1.0, dpi / 72.0)
    matrix = pymupdf.Matrix(zoom, zoom)

    with pymupdf.open(str(src)) as doc:
        page_count = len(doc)
        if page_count == 1:
            out_paths = [
                resolve_output_path(
                    input_path=src,
                    output_dir=output_dir,
                    keep_original_name=keep_original_name,
                    custom_name=custom_name,
                    target_extension=ext,
                    default_suffix="_pagina_01",
                )
            ]
        else:
            out_paths = resolve_multi_output_paths(
                input_path=src,
                part_count=page_count,
                output_dir=output_dir,
                keep_original_name=keep_original_name,
                custom_base_name=custom_name,
                target_extension=ext,
            )

        for idx, page in enumerate(doc):
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            pix.save(str(out_paths[idx]))

    return ConversionFromPDFResult(
        input_path=src,
        output_paths=out_paths,
        target_format=fmt,
        message=f"{len(out_paths)} imagem(ns) {fmt.upper()} gerada(s) a partir do PDF.",
    )

